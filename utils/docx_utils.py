"""DOCX rendering utilities — bold-run helper and resume/cover-letter writer."""
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import config


# ── Paragraph style names used in generated output ────────────────────────────
STYLE_NAME      = "resume_name"
STYLE_HEADER    = "resume_header"
STYLE_JOB_TITLE = "resume_job_title"
STYLE_BULLET    = "resume_bullet"
STYLE_NORMAL    = "resume_normal"
STYLE_CLOSING    = "resume_closing"
STYLE_SIGNATURE  = "resume_signature"
STYLE_SALUTATION = "resume_salutation"

# Mapping from Claude output tag → internal style name
STYLE_MAP = {
    "name":      STYLE_NAME,
    "header":    STYLE_HEADER,
    "job_title": STYLE_JOB_TITLE,
    "bullet":    STYLE_BULLET,
    "normal":    STYLE_NORMAL,
    # cover letter specific
    "date":      STYLE_NORMAL,
    "address":   STYLE_NORMAL,
    "salutation": STYLE_SALUTATION,
    "closing":   STYLE_CLOSING,
    "signature": STYLE_SIGNATURE,
}


def _hex_to_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_runs_with_bold(para, text: str) -> None:
    """Parse **bold** markers into real bold runs on a paragraph."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        elif part:
            para.add_run(part)


def _add_right_tab(para, page_width_inches: float, margin_inches: float) -> None:
    """Add a right-aligned tab stop at the text area right edge."""
    from docx.oxml import OxmlElement as _el
    from docx.oxml.ns import qn as _qn
    text_width = page_width_inches - 2 * margin_inches
    tab_pos_twips = int(text_width * 1440)   # 1 inch = 1440 twips
    pPr = para._p.get_or_add_pPr()
    tabs_el = pPr.find(_qn("w:tabs"))
    if tabs_el is None:
        tabs_el = _el("w:tabs")
        pPr.append(tabs_el)
    tab = _el("w:tab")
    tab.set(_qn("w:val"), "right")
    tab.set(_qn("w:pos"), str(tab_pos_twips))
    tabs_el.append(tab)


def _add_bottom_rule(para, color_hex: str = config.DOCX_NAME_COLOR) -> None:
    """Add a thin horizontal rule below a paragraph via OOXML."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(bottom)
    pPr.append(pBdr)


def _define_styles(doc: Document) -> None:
    """Register custom paragraph styles matching the inferred resume spec."""
    navy = _hex_to_rgb(config.DOCX_NAME_COLOR)
    black = _hex_to_rgb(config.DOCX_BODY_COLOR)
    font = config.DOCX_FONT_FAMILY

    def _base(name: str):
        styles = doc.styles
        try:
            return styles[name]
        except KeyError:
            new_style = styles.add_style(name, 1)   # 1 = WD_STYLE_TYPE.PARAGRAPH
            return new_style

    # Name line — 17pt bold navy
    s = _base(STYLE_NAME)
    s.font.name = font
    s.font.size = Pt(config.DOCX_NAME_SIZE_PT)
    s.font.bold = True
    s.font.color.rgb = navy
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after  = Pt(2)

    # ALL-CAPS section header — 11pt bold navy + bottom rule applied per-paragraph
    s = _base(STYLE_HEADER)
    s.font.name = font
    s.font.size = Pt(config.DOCX_HEADER_SIZE_PT)
    s.font.bold = True
    s.font.color.rgb = navy
    s.paragraph_format.space_before = Pt(6)
    s.paragraph_format.space_after  = Pt(2)
    s.paragraph_format.keep_with_next = True

    # Job title | Company | Dates — 10.5pt bold black
    s = _base(STYLE_JOB_TITLE)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = True
    s.font.color.rgb = black
    s.paragraph_format.space_before = Pt(4)
    s.paragraph_format.space_after  = Pt(1)
    s.paragraph_format.keep_with_next = True

    # Bullet — 10.5pt normal black, indented
    s = _base(STYLE_BULLET)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = False
    s.font.color.rgb = black
    s.paragraph_format.space_before    = Pt(0)
    s.paragraph_format.space_after     = Pt(2)
    s.paragraph_format.left_indent     = Inches(0.25)

    # Normal — contact info, date lines, cover letter body
    s = _base(STYLE_NORMAL)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = False
    s.font.color.rgb = black
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after  = Pt(2)

    # Salutation ("Dear Hiring Manager,") — one blank line above
    s = _base(STYLE_SALUTATION)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = False
    s.font.color.rgb = black
    s.paragraph_format.space_before = Pt(12)   # ~1 line
    s.paragraph_format.space_after  = Pt(2)

    # Closing line ("Sincerely,") — 2 blank lines of space above
    s = _base(STYLE_CLOSING)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = False
    s.font.color.rgb = black
    s.paragraph_format.space_before = Pt(24)   # ~2 lines
    s.paragraph_format.space_after  = Pt(2)

    # Signature — 10.5pt bold navy (smaller than resume name header)
    s = _base(STYLE_SIGNATURE)
    s.font.name = font
    s.font.size = Pt(config.DOCX_BODY_SIZE_PT)
    s.font.bold = True
    s.font.color.rgb = navy
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after  = Pt(2)


def build_document(paragraphs: list[dict], output_path: str) -> None:
    """
    Render a list of {style, text} paragraph dicts to a DOCX file.

    Expected style values (from Claude tool_use output):
      "name"      — candidate name
      "header"    — ALL-CAPS section header
      "job_title" — job title | company | dates
      "bullet"    — achievement bullet
      "normal"    — contact info, date lines, cover letter body
    """
    # Standard US letter: 8.5" wide
    PAGE_WIDTH = 8.5

    doc = Document()

    # Page margins
    m = Inches(config.DOCX_MARGIN_INCHES)
    for sec in doc.sections:
        sec.top_margin    = m
        sec.bottom_margin = m
        sec.left_margin   = m
        sec.right_margin  = m

    # Remove the default blank paragraph Word inserts
    for p in doc.paragraphs:
        p._element.getparent().remove(p._element)

    _define_styles(doc)

    for item in paragraphs:
        tag  = item.get("style", "normal")
        text = item.get("text", "").strip()

        if tag == "page_break":
            para = doc.add_paragraph()
            run  = para.add_run()
            run.add_break(WD_BREAK.PAGE)
            continue

        if not text:
            continue
        style_name = STYLE_MAP.get(tag, STYLE_NORMAL)
        para = doc.add_paragraph(style=style_name)

        if tag == "job_title":
            # Expect text like "Senior Scientist, Bioinformatics | BioNTech | 03/2023 – 09/2025"
            # Split on last " | " to separate title+company from date, render with right tab.
            # If Claude embeds a literal tab already, use it directly.
            if "\t" in text:
                title_part, date_part = text.split("\t", 1)
            elif " | " in text:
                parts = text.rsplit(" | ", 1)
                title_part = parts[0]
                date_part  = parts[1] if len(parts) == 2 else ""
            else:
                title_part = text
                date_part  = ""
            _add_right_tab(para, PAGE_WIDTH, config.DOCX_MARGIN_INCHES)
            if date_part:
                run1 = para.add_run(title_part)
                run1.bold = True
                para.add_run("\t")
                run2 = para.add_run(date_part)
                run2.bold = False
            else:
                _add_runs_with_bold(para, title_part)

        elif tag == "bullet":
            clean = text.lstrip("-– ").lstrip("• ")
            _add_runs_with_bold(para, "• " + clean)

        else:
            _add_runs_with_bold(para, text)

        # Section headers get a bottom rule
        if tag == "header":
            _add_bottom_rule(para)

    doc.save(output_path)
    print(f"  Saved → {output_path}")
