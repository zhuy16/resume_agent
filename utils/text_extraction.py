"""Consolidated text and structure extraction from PDF and DOCX files."""
import os

from PyPDF2 import PdfReader
from docx import Document


def extract_text_from_pdf(path: str, max_chars: int = 8000) -> str:
    """Extract plain text from a PDF, up to max_chars."""
    try:
        reader = PdfReader(path)
        text = ""
        for page in reader.pages:
            text += (page.extract_text() or "") + " "
            if len(text) >= max_chars:
                break
        return text[:max_chars].strip()
    except Exception as e:
        print(f"  [warn] Could not read PDF {os.path.basename(path)}: {e}")
        return ""


def extract_text_from_docx(path: str, max_chars: int = 8000) -> str:
    """Extract plain text from a DOCX, up to max_chars."""
    try:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return text[:max_chars]
    except Exception as e:
        print(f"  [warn] Could not read DOCX {os.path.basename(path)}: {e}")
        return ""


def extract_text(path: str, max_chars: int = 8000) -> str:
    """Extract plain text from PDF or DOCX based on extension."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path, max_chars)
    if ext == ".docx":
        return extract_text_from_docx(path, max_chars)
    try:
        with open(path, encoding="utf-8") as f:
            return f.read(max_chars)
    except Exception as e:
        print(f"  [warn] Could not read {os.path.basename(path)}: {e}")
        return ""


def extract_structured_paragraphs(path: str) -> list[dict]:
    """
    Extract paragraphs from a DOCX as a list of {style, text} dicts,
    preserving the paragraph structure for faithful round-trip to Claude.
    Falls back to flat lines for PDF.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            doc = Document(path)
            return [
                {"style": p.style.name, "text": p.text}
                for p in doc.paragraphs
                if p.text.strip()
            ]
        except Exception as e:
            print(f"  [warn] Could not read structured DOCX {os.path.basename(path)}: {e}")
            return []
    else:
        raw = extract_text_from_pdf(path)
        return [
            {"style": "Normal", "text": line}
            for line in raw.splitlines()
            if line.strip()
        ]
